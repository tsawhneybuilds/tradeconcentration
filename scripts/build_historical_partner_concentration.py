#!/usr/bin/env python3
"""Build the long-run TRADHIST partner concentration panel."""

from __future__ import annotations

import argparse
import math
from collections import Counter
from pathlib import Path
from typing import Any

import numpy as np
import pandas as pd
import requests
from openpyxl import load_workbook

from concentration_metrics import (
    active_effective_count,
    active_gini,
    active_hhi,
    active_normalized_gini,
    active_normalized_hhi,
    active_theil,
    active_top_share,
)
from historical_partner_common import (
    ENTITY_SPECS,
    EXPECTED_ENTITY_INTERIOR_GAPS,
    EXPECTED_ENTITY_SPANS,
    FIGURES_DIR,
    FLOW_EXPORTS,
    FLOW_IMPORTS,
    INVALID_PARTNER_CODES,
    MPD_SOURCE_URLS,
    PRIMARY_PARTNER_THRESHOLD,
    PROCESSED_DIR,
    RAW_DIR,
    RESULTS_DIR,
    SELECTED_ENTITY_ORDER,
    TRADHIST_END_YEAR,
    TRADHIST_EXPECTED_ROWS,
    TRADHIST_SOURCE_URLS,
    TRADHIST_START_YEAR,
    boundary_flag,
    ensure_dirs,
    entity_metadata_rows,
    finite_nonnegative,
    json_default,
    mpd_code_for_entity,
    now_utc,
    primary_source_paths,
    mpd_source_paths,
    quality_band,
    relpath,
    source_family,
    source_file_manifest,
    supports_headline,
    write_json,
)

EXPECTED_TRADE_COLUMNS = ["iso_o", "iso_d", "year", "FLOW", "FLOW_0", "SOURCE_TF"]
OPTIONAL_TRADE_COLUMNS = ["BITARIFF"]
EXPECTED_YEAR_UNIVERSE_COLUMNS = ["iso", "year"]
EXPECTED_COUNTRY_SPECIFIC_COLUMNS = ["iso", "CONTI", "REGIO"]
SELECTED_ENTITY_SET = set(SELECTED_ENTITY_ORDER)
CHUNK_BYTES = 1 << 20


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--skip-download", action="store_true", help="Use existing raw files and fail if missing.")
    parser.add_argument("--force-redownload", action="store_true", help="Refresh raw downloads even if files exist.")
    return parser.parse_args()


def download_file(url: str, destination: Path, force: bool = False) -> None:
    destination.parent.mkdir(parents=True, exist_ok=True)
    if destination.exists() and not force:
        return
    response = requests.get(url, stream=True, timeout=120)
    response.raise_for_status()
    with destination.open("wb") as handle:
        for chunk in response.iter_content(chunk_size=CHUNK_BYTES):
            if chunk:
                handle.write(chunk)


def ensure_raw_sources(skip_download: bool, force_redownload: bool) -> dict[str, list[dict[str, Any]]]:
    ensure_dirs()
    manifests: dict[str, list[dict[str, Any]]] = {"tradhist": [], "maddison": []}
    for name, url in TRADHIST_SOURCE_URLS.items():
        path = primary_source_paths()[name]
        if not skip_download:
            download_file(url, path, force=force_redownload)
        if not path.exists():
            raise FileNotFoundError(f"Missing required TRADHIST source: {path}")
        manifests["tradhist"].append(source_file_manifest(path, url))
    for name, url in MPD_SOURCE_URLS.items():
        path = mpd_source_paths()[name]
        if not skip_download:
            download_file(url, path, force=force_redownload)
        if not path.exists():
            raise FileNotFoundError(f"Missing required Maddison source: {path}")
        manifests["maddison"].append(source_file_manifest(path, url))
    source_registry = {
        "generated_at_utc": now_utc(),
        "tradhist": manifests["tradhist"],
        "maddison": manifests["maddison"],
        "citation": {
            "tradhist": "Fouquin and Hugot (2016), CEPII TRADHIST working paper and bilateral workbook release.",
            "maddison": "Bolt and Van Zanden (2024), Maddison Project Database 2023 release.",
        },
    }
    write_json(RAW_DIR / "source_registry.json", source_registry)
    return manifests


def _read_openpyxl_rows(path: Path) -> tuple[list[str], Any]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)
    header = [str(value) for value in next(rows)]
    return header, (workbook, rows)


def read_valid_entities() -> tuple[set[str], dict[str, dict[str, Any]], dict[int, set[str]]]:
    country_path = primary_source_paths()["TRADHIST_GRAVITY_COUNTRY_SPECIFIC.xlsx"]
    year_path = primary_source_paths()["TRADHIST_GRAVITY_COUNTRY_YEAR_SPECIFIC.xlsx"]

    header, state = _read_openpyxl_rows(country_path)
    workbook, rows = state
    idx = {name: pos for pos, name in enumerate(header)}
    missing = [name for name in EXPECTED_COUNTRY_SPECIFIC_COLUMNS if name not in idx]
    if missing:
        raise RuntimeError(f"Country-specific TRADHIST workbook is missing columns: {missing}")
    valid_codes: set[str] = set()
    metadata: dict[str, dict[str, Any]] = {}
    for row in rows:
        code = str(row[idx["iso"]]).strip()
        if not code:
            continue
        valid_codes.add(code)
        metadata[code] = {
            "iso": code,
            "continent": row[idx["CONTI"]],
            "region": row[idx["REGIO"]],
        }
    workbook.close()

    header, state = _read_openpyxl_rows(year_path)
    workbook, rows = state
    idx = {name: pos for pos, name in enumerate(header)}
    missing = [name for name in EXPECTED_YEAR_UNIVERSE_COLUMNS if name not in idx]
    if missing:
        raise RuntimeError(f"Country-year TRADHIST workbook is missing columns: {missing}")
    valid_by_year: dict[int, set[str]] = {}
    for row in rows:
        code_raw = row[idx["iso"]]
        year_raw = row[idx["year"]]
        if code_raw is None or year_raw is None:
            continue
        code = str(code_raw).strip()
        year = int(year_raw)
        valid_by_year.setdefault(year, set()).add(code)
    workbook.close()
    return valid_codes, metadata, valid_by_year


def _append_role_record(
    records: list[dict[str, Any]],
    *,
    reporter: str,
    partner: str,
    flow: str,
    year: int,
    flow_value: float | None,
    flow0_value: float | None,
    source_tf: str,
    bitariff: float | None,
    source_file: str,
    part_label: str,
    valid_entity_codes: set[str],
    entity_metadata: dict[str, dict[str, Any]],
) -> None:
    partner_code = partner.strip()
    reporter_code = reporter.strip()
    spec = ENTITY_SPECS[reporter_code]
    partner_valid = partner_code not in INVALID_PARTNER_CODES and partner_code in valid_entity_codes
    observed_zero = flow_value == 0.0 if flow_value is not None else False
    likely_zero = flow_value is None and flow0_value == 0.0
    records.append(
        {
            "entity_id": reporter_code,
            "entity_label": spec.entity_label,
            "entity_boundary_note": spec.boundary_note,
            "flow": flow,
            "partner_id": partner_code,
            "year": year,
            "FLOW": flow_value,
            "FLOW_0": flow0_value,
            "SOURCE_TF": source_tf,
            "source_family": source_family(source_tf),
            "BITARIFF": bitariff,
            "source_file": source_file,
            "source_part": part_label,
            "is_self_trade": reporter_code == partner_code,
            "partner_in_master_codebook": partner_valid,
            "partner_included_main_panel": partner_valid and reporter_code != partner_code,
            "observed_positive": flow_value is not None and flow_value > 0,
            "observed_zero": observed_zero,
            "likely_zero": likely_zero,
            "coded_zero": observed_zero or likely_zero,
            "value_current_gbp": flow_value if flow_value is not None else 0.0 if observed_zero else np.nan,
            "partner_continent": entity_metadata.get(partner_code, {}).get("continent"),
            "partner_region": entity_metadata.get(partner_code, {}).get("region"),
            "boundary_flag": boundary_flag(reporter_code, year),
            "reporter_role": "origin" if flow == FLOW_EXPORTS else "destination",
        }
    )


def build_partner_flows(
    valid_entity_codes: set[str],
    entity_metadata: dict[str, dict[str, Any]],
) -> tuple[pd.DataFrame, dict[str, Any]]:
    trade_paths = [
        primary_source_paths()["TRADHIST_BITRADE_BITARIFF_1.xlsx"],
        primary_source_paths()["TRADHIST_BITRADE_BITARIFF_2.xlsx"],
        primary_source_paths()["TRADHIST_BITRADE_BITARIFF_3.xlsx"],
    ]
    selected_records: list[dict[str, Any]] = []
    total_rows = 0
    selected_rows = 0
    part_rows: dict[str, int] = {}
    part_selected_rows: dict[str, int] = {}
    source_counter: Counter[str] = Counter()

    for part_no, path in enumerate(trade_paths, start=1):
        header, state = _read_openpyxl_rows(path)
        workbook, rows = state
        idx = {name: pos for pos, name in enumerate(header)}
        missing = [name for name in EXPECTED_TRADE_COLUMNS if name not in idx]
        if missing:
            raise RuntimeError(f"Trade workbook {path.name} is missing columns: {missing}")
        part_label = f"part_{part_no}"
        part_rows[part_label] = 0
        part_selected_rows[part_label] = 0
        for row in rows:
            part_rows[part_label] += 1
            total_rows += 1
            iso_o = str(row[idx["iso_o"]]).strip()
            iso_d = str(row[idx["iso_d"]]).strip()
            year = int(row[idx["year"]])
            flow_value = finite_nonnegative(row[idx["FLOW"]])
            flow0_value = finite_nonnegative(row[idx["FLOW_0"]])
            source_tf = str(row[idx["SOURCE_TF"]] or "").strip()
            bitariff = finite_nonnegative(row[idx["BITARIFF"]]) if "BITARIFF" in idx else None

            if iso_o not in SELECTED_ENTITY_SET and iso_d not in SELECTED_ENTITY_SET:
                continue

            selected_rows += 1
            part_selected_rows[part_label] += 1
            source_counter[source_tf] += 1

            if iso_o in SELECTED_ENTITY_SET:
                _append_role_record(
                    selected_records,
                    reporter=iso_o,
                    partner=iso_d,
                    flow=FLOW_EXPORTS,
                    year=year,
                    flow_value=flow_value,
                    flow0_value=flow0_value,
                    source_tf=source_tf,
                    bitariff=bitariff,
                    source_file=path.name,
                    part_label=part_label,
                    valid_entity_codes=valid_entity_codes,
                    entity_metadata=entity_metadata,
                )
            if iso_d in SELECTED_ENTITY_SET:
                _append_role_record(
                    selected_records,
                    reporter=iso_d,
                    partner=iso_o,
                    flow=FLOW_IMPORTS,
                    year=year,
                    flow_value=flow_value,
                    flow0_value=flow0_value,
                    source_tf=source_tf,
                    bitariff=bitariff,
                    source_file=path.name,
                    part_label=part_label,
                    valid_entity_codes=valid_entity_codes,
                    entity_metadata=entity_metadata,
                )
        workbook.close()

    if total_rows != TRADHIST_EXPECTED_ROWS:
        raise RuntimeError(f"TRADHIST combined row count mismatch: expected {TRADHIST_EXPECTED_ROWS}, found {total_rows}")

    flows = pd.DataFrame.from_records(selected_records)
    flows["year"] = flows["year"].astype(int)
    duplicate_mask = flows.duplicated(["entity_id", "partner_id", "year", "flow"], keep=False)
    if duplicate_mask.any():
        dupes = flows.loc[duplicate_mask, ["entity_id", "partner_id", "year", "flow", "SOURCE_TF"]].copy()
        dup_path = RESULTS_DIR / "historical_partner_concentration_duplicate_dyads.csv"
        dupes.to_csv(dup_path, index=False)
        examples = dupes.head(10).to_dict(orient="records")
        raise RuntimeError(f"Selected TRADHIST dyads are not unique; see {dup_path}. Examples: {examples}")

    diagnostics = {
        "tradhist_total_rows": total_rows,
        "selected_source_rows": selected_rows,
        "selected_partner_flow_rows": int(len(flows)),
        "trade_rows_by_part": part_rows,
        "selected_rows_by_part": part_selected_rows,
        "top_selected_sources": source_counter.most_common(25),
    }
    return flows.sort_values(["entity_id", "year", "flow", "partner_id"]).reset_index(drop=True), diagnostics


def _zero_inclusive_gini(values: np.ndarray, universe_count: int) -> float:
    arr = np.asarray(values, dtype=float)
    arr = arr[np.isfinite(arr) & (arr >= 0)]
    if universe_count <= 0 or arr.size == 0 or float(arr.sum()) <= 0:
        return np.nan
    positive = arr[arr > 0]
    if positive.size == 0:
        return np.nan
    padded = np.concatenate([positive, np.zeros(max(0, universe_count - positive.size), dtype=float)])
    padded.sort()
    total = float(padded.sum())
    n = padded.size
    ranks = np.arange(1, n + 1, dtype=float)
    return float((2 * np.sum(ranks * padded) / (n * total)) - ((n + 1) / n))


def _metric_row(
    all_rows: pd.DataFrame,
    included_rows: pd.DataFrame,
    year_universe: dict[int, set[str]],
) -> dict[str, Any]:
    entity_id = str(all_rows["entity_id"].iloc[0])
    year = int(all_rows["year"].iloc[0])
    flow = str(all_rows["flow"].iloc[0])
    spec = ENTITY_SPECS[entity_id]

    universe = set(year_universe.get(year, set()))
    universe_count = len(universe - {entity_id}) if entity_id in universe else len(universe)
    positive_rows = included_rows[included_rows["observed_positive"].astype(bool)].copy()
    values = positive_rows["FLOW"].to_numpy(dtype=float) if not positive_rows.empty else np.array([], dtype=float)
    active_partner_count = int(len(positive_rows))
    observed_zero_count = int(included_rows["observed_zero"].astype(bool).sum())
    likely_zero_count = int(included_rows["likely_zero"].astype(bool).sum())
    coded_zero_count = observed_zero_count + likely_zero_count
    explicit_missing_rows = int(
        (
            (~included_rows["observed_positive"].astype(bool))
            & (~included_rows["observed_zero"].astype(bool))
            & (~included_rows["likely_zero"].astype(bool))
        ).sum()
    )
    missing_unobserved_count = max(0, universe_count - active_partner_count - coded_zero_count)
    raw_total = float(np.nansum(included_rows["FLOW"].to_numpy(dtype=float))) if not included_rows.empty else 0.0

    by_source = positive_rows.groupby("SOURCE_TF", dropna=False).agg(
        positive_value_current_gbp=("FLOW", "sum"),
        active_partner_count=("partner_id", "nunique"),
    )
    dominant_source_tf = ""
    dominant_source_value_share = np.nan
    dominant_source_partner_share = np.nan
    dominant_source_family = "Unknown"
    if not by_source.empty:
        dominant_source_tf = str(by_source["positive_value_current_gbp"].idxmax() or "")
        dominant_source_family = source_family(dominant_source_tf)
        total_positive = float(by_source["positive_value_current_gbp"].sum())
        if total_positive > 0:
            dominant_source_value_share = float(
                by_source.loc[dominant_source_tf, "positive_value_current_gbp"] / total_positive
            )
        if active_partner_count > 0:
            dominant_source_partner_share = float(by_source.loc[dominant_source_tf, "active_partner_count"] / active_partner_count)
    elif not included_rows.empty:
        dominant_source_tf = str(included_rows["SOURCE_TF"].mode(dropna=False).iloc[0] or "")
        dominant_source_family = source_family(dominant_source_tf)

    source_family_stats = positive_rows.groupby("source_family", dropna=False).agg(
        positive_value_current_gbp=("FLOW", "sum"),
        active_partner_count=("partner_id", "nunique"),
    )
    rows: dict[str, Any] = {
        "entity_id": entity_id,
        "entity_label": spec.entity_label,
        "entity_boundary_note": spec.boundary_note,
        "flow": flow,
        "year": year,
        "boundary_flag": boundary_flag(entity_id, year),
        "bilateral_sum_current_gbp": raw_total,
        "active_partner_count": active_partner_count,
        "explicit_observed_zero_count": observed_zero_count,
        "likely_zero_count": likely_zero_count,
        "coded_zero_count": coded_zero_count,
        "missing_unobserved_count": missing_unobserved_count,
        "explicit_missing_observed_rows": explicit_missing_rows,
        "universe_partner_count": universe_count,
        "quality_band": quality_band(active_partner_count),
        "headline_eligible": supports_headline(active_partner_count),
        "appendix_eligible": active_partner_count >= 10,
        "diagnostic_only": 5 <= active_partner_count < 10,
        "metrics_suppressed": active_partner_count < 5,
        "dominant_source_tf": dominant_source_tf,
        "dominant_source_family": dominant_source_family,
        "dominant_source_value_share": dominant_source_value_share,
        "dominant_source_partner_share": dominant_source_partner_share,
        "year_is_wwi": 1914 <= year <= 1918,
        "year_is_wwii": 1939 <= year <= 1945,
        "year_post_1948": year >= 1948,
        "year_post_1991": year >= 1991,
        "source_regime_flag": "post1948_dots" if year >= 1948 and dominant_source_family == "DOTS" else (
            "post1948_other" if year >= 1948 else "pre1948"
        ),
    }

    if active_partner_count >= 1 and np.isfinite(values).any():
        rows["top1_partner_share"] = active_top_share(values, n=1)
        rows["top3_partner_share"] = active_top_share(values, n=3)
        rows["top5_partner_share"] = active_top_share(values, n=5)
        rows["top10_partner_share"] = active_top_share(values, n=10)
        rows["partner_hhi"] = active_hhi(values)
        rows["effective_partner_count"] = active_effective_count(values)
    else:
        rows["top1_partner_share"] = np.nan
        rows["top3_partner_share"] = np.nan
        rows["top5_partner_share"] = np.nan
        rows["top10_partner_share"] = np.nan
        rows["partner_hhi"] = np.nan
        rows["effective_partner_count"] = np.nan

    if active_partner_count >= 5:
        rows["partner_gini"] = active_gini(values)
        rows["partner_gini_normalized"] = active_normalized_gini(values)
        rows["partner_theil"] = active_theil(values, normalized=False)
        rows["partner_theil_normalized"] = active_theil(values, normalized=True)
        rows["partner_hhi_normalized"] = active_normalized_hhi(values)
    else:
        rows["partner_gini"] = np.nan
        rows["partner_gini_normalized"] = np.nan
        rows["partner_theil"] = np.nan
        rows["partner_theil_normalized"] = np.nan
        rows["partner_hhi_normalized"] = np.nan

    rows["coded_zero_inclusive_partner_gini"] = _zero_inclusive_gini(values, active_partner_count + coded_zero_count)

    family_payload = []
    total_positive_value = float(source_family_stats["positive_value_current_gbp"].sum()) if not source_family_stats.empty else 0.0
    for family, stats in source_family_stats.reset_index().iterrows():
        family_payload.append(
            {
                "entity_id": entity_id,
                "flow": flow,
                "year": year,
                "source_family": stats["source_family"],
                "positive_value_current_gbp": float(stats["positive_value_current_gbp"]),
                "active_partner_count": int(stats["active_partner_count"]),
                "positive_value_share": float(stats["positive_value_current_gbp"] / total_positive_value)
                if total_positive_value > 0
                else np.nan,
                "active_partner_share": float(stats["active_partner_count"] / active_partner_count)
                if active_partner_count > 0
                else np.nan,
            }
        )
    rows["source_family_payload"] = family_payload
    return rows


def build_flow_panel(flows: pd.DataFrame, year_universe: dict[int, set[str]]) -> tuple[pd.DataFrame, pd.DataFrame]:
    metrics_rows: list[dict[str, Any]] = []
    diagnostics_rows: list[dict[str, Any]] = []
    grouped_all = flows.groupby(["entity_id", "year", "flow"], sort=True)
    for _, all_rows in grouped_all:
        included_rows = all_rows[all_rows["partner_included_main_panel"].astype(bool)].copy()
        row = _metric_row(all_rows, included_rows, year_universe)
        diagnostics_rows.extend(row.pop("source_family_payload"))
        metrics_rows.append(row)
    panel = pd.DataFrame.from_records(metrics_rows).sort_values(["entity_id", "year", "flow"]).reset_index(drop=True)
    diagnostics = pd.DataFrame.from_records(diagnostics_rows).sort_values(
        ["entity_id", "year", "flow", "source_family"]
    ).reset_index(drop=True)
    return panel, diagnostics


def read_maddison_controls() -> pd.DataFrame:
    path = mpd_source_paths()["mpd2023_web.xlsx"]
    mpd = pd.read_excel(path, sheet_name="Full data")
    required = {"countrycode", "country", "region", "year", "gdppc", "pop"}
    missing = sorted(required - set(mpd.columns))
    if missing:
        raise RuntimeError(f"Maddison full-data sheet is missing columns: {missing}")
    mpd = mpd[list(required)].copy()
    mpd["countrycode"] = mpd["countrycode"].astype(str).str.upper().str.strip()
    mpd["year"] = pd.to_numeric(mpd["year"], errors="coerce").astype("Int64")
    mpd["gdppc"] = pd.to_numeric(mpd["gdppc"], errors="coerce")
    mpd["pop"] = pd.to_numeric(mpd["pop"], errors="coerce")
    mpd = mpd.dropna(subset=["year"]).copy()
    mpd["year"] = mpd["year"].astype(int)
    return mpd


def merge_maddison_controls(panel: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame]:
    mpd = read_maddison_controls()
    merged = panel.copy()
    merged["mpd_countrycode"] = [
        mpd_code_for_entity(entity_id, int(year)) for entity_id, year in zip(merged["entity_id"], merged["year"], strict=True)
    ]
    merged = merged.merge(
        mpd.rename(
            columns={
                "countrycode": "mpd_countrycode",
                "country": "mpd_country",
                "region": "mpd_region",
                "gdppc": "mpd_gdppc_2011_usd",
                "pop": "mpd_population_thousands",
            }
        ),
        on=["mpd_countrycode", "year"],
        how="left",
        validate="many_to_one",
    )
    merged["gdppc_10k"] = merged["mpd_gdppc_2011_usd"] / 10_000.0
    merged["gdppc_10k_sq"] = merged["gdppc_10k"] ** 2
    merged["log_gdppc"] = np.where(merged["mpd_gdppc_2011_usd"] > 0, np.log(merged["mpd_gdppc_2011_usd"]), np.nan)
    merged["log_gdppc_sq"] = merged["log_gdppc"] ** 2
    merged["log_population"] = np.where(
        merged["mpd_population_thousands"] > 0, np.log(merged["mpd_population_thousands"]), np.nan
    )
    attrition = (
        merged.assign(
            mpd_gdppc_missing=merged["mpd_gdppc_2011_usd"].isna(),
            mpd_population_missing=merged["mpd_population_thousands"].isna(),
        )[
            [
                "entity_id",
                "year",
                "flow",
                "mpd_countrycode",
                "mpd_gdppc_missing",
                "mpd_population_missing",
            ]
        ]
        .sort_values(["entity_id", "year", "flow"])
        .reset_index(drop=True)
    )
    return merged, attrition


def build_wide_panel(flow_panel: pd.DataFrame) -> pd.DataFrame:
    common_cols = [
        "entity_id",
        "entity_label",
        "entity_boundary_note",
        "year",
        "mpd_countrycode",
        "mpd_country",
        "mpd_region",
        "mpd_gdppc_2011_usd",
        "mpd_population_thousands",
        "gdppc_10k",
        "gdppc_10k_sq",
        "log_gdppc",
        "log_gdppc_sq",
        "log_population",
    ]
    flow_specific = [
        "bilateral_sum_current_gbp",
        "active_partner_count",
        "explicit_observed_zero_count",
        "likely_zero_count",
        "coded_zero_count",
        "missing_unobserved_count",
        "explicit_missing_observed_rows",
        "universe_partner_count",
        "quality_band",
        "headline_eligible",
        "appendix_eligible",
        "diagnostic_only",
        "metrics_suppressed",
        "dominant_source_tf",
        "dominant_source_family",
        "dominant_source_value_share",
        "dominant_source_partner_share",
        "partner_gini",
        "partner_gini_normalized",
        "partner_theil",
        "partner_theil_normalized",
        "partner_hhi",
        "partner_hhi_normalized",
        "effective_partner_count",
        "top1_partner_share",
        "top3_partner_share",
        "top5_partner_share",
        "top10_partner_share",
        "coded_zero_inclusive_partner_gini",
        "boundary_flag",
        "year_is_wwi",
        "year_is_wwii",
        "year_post_1948",
        "year_post_1991",
        "source_regime_flag",
    ]
    export_panel = flow_panel[flow_panel["flow"] == FLOW_EXPORTS][common_cols + flow_specific].copy()
    import_panel = flow_panel[flow_panel["flow"] == FLOW_IMPORTS][common_cols + flow_specific].copy()
    export_panel = export_panel.rename(
        columns={
            **{col: f"export_{col}" for col in flow_specific if col != "bilateral_sum_current_gbp"},
            "bilateral_sum_current_gbp": "bilateral_sum_exports_current_gbp",
        }
    )
    import_panel = import_panel.rename(
        columns={
            **{col: f"import_{col}" for col in flow_specific if col != "bilateral_sum_current_gbp"},
            "bilateral_sum_current_gbp": "bilateral_sum_imports_current_gbp",
        }
    )
    shared_keys = ["entity_id", "entity_label", "entity_boundary_note", "year"]
    right_shared = [col for col in common_cols if col not in shared_keys]
    wide = export_panel.merge(
        import_panel.drop(columns=right_shared),
        on=shared_keys,
        how="outer",
        validate="one_to_one",
    )
    return wide.sort_values(["entity_id", "year"]).reset_index(drop=True)


def validate_entity_spans(flow_panel: pd.DataFrame) -> None:
    failures = []
    for entity_id, (expected_first, expected_last) in EXPECTED_ENTITY_SPANS.items():
        entity_years = sorted(flow_panel.loc[flow_panel["entity_id"].eq(entity_id), "year"].unique().tolist())
        if not entity_years:
            failures.append({"entity_id": entity_id, "problem": "missing from flow panel"})
            continue
        actual_first = int(entity_years[0])
        actual_last = int(entity_years[-1])
        if actual_first != expected_first or actual_last != expected_last:
            failures.append(
                {
                    "entity_id": entity_id,
                    "expected_first_year": expected_first,
                    "actual_first_year": actual_first,
                    "expected_last_year": expected_last,
                    "actual_last_year": actual_last,
                }
            )
            continue
        missing_years = sorted(set(range(expected_first, expected_last + 1)) - set(entity_years))
        allowed_gaps = sorted(EXPECTED_ENTITY_INTERIOR_GAPS.get(entity_id, set()))
        if missing_years != allowed_gaps:
            failures.append(
                {
                    "entity_id": entity_id,
                    "expected_interior_gaps": allowed_gaps,
                    "actual_missing_years": missing_years,
                }
            )
    if failures:
        raise RuntimeError(f"Entity span validation failed: {failures[:5]}")


def save_outputs(
    partner_flows: pd.DataFrame,
    flow_panel: pd.DataFrame,
    wide_panel: pd.DataFrame,
    source_diagnostics: pd.DataFrame,
    merge_attrition: pd.DataFrame,
    source_manifests: dict[str, list[dict[str, Any]]],
    build_diagnostics: dict[str, Any],
) -> dict[str, Any]:
    PROCESSED_DIR.mkdir(parents=True, exist_ok=True)
    RESULTS_DIR.mkdir(parents=True, exist_ok=True)
    FIGURES_DIR.mkdir(parents=True, exist_ok=True)

    flow_path = PROCESSED_DIR / "tradhist_partner_flows.parquet"
    long_panel_path = PROCESSED_DIR / "historical_partner_concentration_long.parquet"
    wide_panel_path = PROCESSED_DIR / "historical_partner_concentration_panel.parquet"
    csv_panel_path = RESULTS_DIR / "historical_partner_concentration_panel.csv"
    source_diag_path = RESULTS_DIR / "source_coverage_diagnostics.csv"
    merge_attrition_path = RESULTS_DIR / "historical_partner_mpd_merge_attrition.csv"
    build_manifest_path = RESULTS_DIR / "historical_partner_build_manifest.json"

    partner_flows.to_parquet(flow_path, index=False)
    flow_panel.to_parquet(long_panel_path, index=False)
    wide_panel.to_parquet(wide_panel_path, index=False)
    wide_panel.to_csv(csv_panel_path, index=False)
    source_diagnostics.to_csv(source_diag_path, index=False)
    merge_attrition.to_csv(merge_attrition_path, index=False)

    build_manifest = {
        "generated_at_utc": now_utc(),
        "source_manifests": source_manifests,
        "build_diagnostics": build_diagnostics,
        "selected_entities": entity_metadata_rows(),
        "output_files": {
            "tradhist_partner_flows_parquet": relpath(flow_path),
            "historical_partner_concentration_long_parquet": relpath(long_panel_path),
            "historical_partner_concentration_panel_parquet": relpath(wide_panel_path),
            "historical_partner_concentration_panel_csv": relpath(csv_panel_path),
            "source_coverage_diagnostics_csv": relpath(source_diag_path),
            "historical_partner_mpd_merge_attrition_csv": relpath(merge_attrition_path),
        },
        "units": {
            "trade_value": "current British pounds",
            "mpd_gdppc": "Real GDP per capita in 2011$",
            "mpd_population": "Population, mid-year (thousands)",
        },
    }
    write_json(build_manifest_path, build_manifest)
    return build_manifest


def main() -> None:
    args = parse_args()
    source_manifests = ensure_raw_sources(args.skip_download, args.force_redownload)
    valid_entity_codes, entity_metadata, year_universe = read_valid_entities()
    partner_flows, build_diagnostics = build_partner_flows(valid_entity_codes, entity_metadata)

    partner_flows["partner_exists_in_year_universe"] = [
        partner in year_universe.get(int(year), set())
        for partner, year in zip(partner_flows["partner_id"], partner_flows["year"], strict=True)
    ]
    partner_flows["partner_included_main_panel"] = (
        partner_flows["partner_included_main_panel"].astype(bool)
        & partner_flows["partner_exists_in_year_universe"].astype(bool)
    )
    flow_panel, source_diagnostics = build_flow_panel(partner_flows, year_universe)
    flow_panel, merge_attrition = merge_maddison_controls(flow_panel)
    validate_entity_spans(flow_panel)
    wide_panel = build_wide_panel(flow_panel)

    manifest = save_outputs(
        partner_flows=partner_flows,
        flow_panel=flow_panel,
        wide_panel=wide_panel,
        source_diagnostics=source_diagnostics,
        merge_attrition=merge_attrition,
        source_manifests=source_manifests,
        build_diagnostics=build_diagnostics,
    )
    print(f"Built historical partner concentration panel at {manifest['output_files']['historical_partner_concentration_panel_parquet']}")


if __name__ == "__main__":
    main()
